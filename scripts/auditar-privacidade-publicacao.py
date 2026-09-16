"""Revalida CPF, telefone e marca de assinatura no lote público reconciliado."""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


RAIZ = Path(__file__).resolve().parents[1]
RECONCILIACAO = RAIZ / "docs/carga/RECONCILIACAO_INDEPENDENTE_2026-09-16.json"


def carregar_ambiente() -> None:
    arquivo = RAIZ / ".env.local"
    if not arquivo.exists():
        return
    for linha in arquivo.read_text(encoding="utf-8").splitlines():
        if not linha or linha.lstrip().startswith("#") or "=" not in linha:
            continue
        chave, valor = linha.split("=", 1)
        os.environ.setdefault(chave, valor.strip().strip('"').strip("'"))


def cpf_valido(digitos: str) -> bool:
    if len(digitos) != 11 or len(set(digitos)) == 1:
        return False
    numeros = [int(valor) for valor in digitos]
    primeiro = 11 - sum(numeros[i] * (10 - i) for i in range(9)) % 11
    primeiro = 0 if primeiro >= 10 else primeiro
    segundo = 11 - sum(numeros[i] * (11 - i) for i in range(10)) % 11
    segundo = 0 if segundo >= 10 else segundo
    return numeros[9:] == [primeiro, segundo]


def texto_do_arquivo(caminho: Path) -> str:
    sufixo = caminho.suffix.lower()
    if sufixo == ".pdf":
        return "\n".join(pagina.extract_text() or "" for pagina in PdfReader(caminho).pages)
    if sufixo == ".xlsx":
        pasta = load_workbook(caminho, read_only=True, data_only=True)
        partes: list[str] = []
        for planilha in pasta.worksheets:
            for linha in planilha.iter_rows(values_only=True):
                partes.extend(str(valor) for valor in linha if valor is not None)
        pasta.close()
        return "\n".join(partes)
    return caminho.read_text(encoding="utf-8")


def main() -> int:
    carregar_ambiente()
    raiz_fontes = Path(os.environ["OBSERVATORIO_FONTES_DIR"]).resolve()
    dados = json.loads(RECONCILIACAO.read_text(encoding="utf-8"))
    origens = {
        item["fonte_original"]
        for item in dados["objetos"]
        if Path(item["fonte_original"]).suffix.lower() in {".pdf", ".xlsx"}
    }
    ocr = raiz_fontes / "derivados/a03-borda-da-mata-ocr.md"
    caminhos = [(origem, raiz_fontes / origem) for origem in sorted(origens)]
    caminhos.append(("derivados/a03-borda-da-mata-ocr.md", ocr))

    cpf_padrao = re.compile(r"(?<!\d)(\d{3}[.\s-]?\d{3}[.\s-]?\d{3}[-.\s]?\d{2})(?!\d)")
    telefone_padrao = re.compile(
        r"(?<![\d.])(?:\+?55\s*)?(?:\(?\d{2}\)?[\s.-]*)?9?\d{4}[\s-]\d{4}(?![\d.])"
    )
    assinatura_padrao = re.compile(r"\bassinatura(?:\s+digital)?\s*[:_]", re.IGNORECASE)
    achados: list[dict[str, str]] = []
    sem_texto: list[str] = []

    for origem, caminho in caminhos:
        texto = texto_do_arquivo(caminho)
        if caminho.suffix.lower() == ".pdf" and not texto.strip():
            sem_texto.append(origem)
        cpfs = [valor for valor in cpf_padrao.findall(texto) if cpf_valido(re.sub(r"\D", "", valor))]
        if cpfs:
            achados.append({"origem": origem, "categoria": "CPF", "quantidade": str(len(cpfs))})
        telefones = telefone_padrao.findall(texto)
        if telefones:
            achados.append(
                {"origem": origem, "categoria": "telefone", "quantidade": str(len(telefones))}
            )
        assinaturas = assinatura_padrao.findall(texto)
        if assinaturas:
            achados.append(
                {"origem": origem, "categoria": "assinatura", "quantidade": str(len(assinaturas))}
            )

    resultado = {
        "arquivos_publicos_textuais_analisados": len(origens),
        "transcricao_visual_a03_analisada": ocr.exists(),
        "pdfs_sem_camada_textual": sem_texto,
        "achados": achados,
    }
    print(json.dumps(resultado, ensure_ascii=False, indent=2))
    return 1 if achados else 0


if __name__ == "__main__":
    sys.exit(main())
